#!/usr/bin/env python3
"""Tapu/Parsel Rapor Üretici — Excel (.xlsx) ve Word (.docx) çıktısı.

Kullanım:
    python tapu_raporu.py --input sonuclar.json --format xlsx
    python tapu_raporu.py --input sonuclar.json --format docx --baslik "Rapor Adı"
    python tapu_raporu.py --input sonuclar.json --format both --tarih 2026-07-15

JSON girdi formatı (parsel_sorgu.py çıktısıyla uyumlu):
    [
      {
        "ada": "123", "parsel": "4", "il": "İstanbul", "ilce": "Kadıköy",
        "mahalle": "Moda", "alan_m2": 1250.5, "nitelik": "Arsa",
        "harita_url": "https://maps.google.com/...",
        "aciklama": "Proje alanı"
      },
      ...
    ]

Çıktı dosyaları: ./out/tapu_raporu.xlsx ve/veya ./out/tapu_raporu.docx
"""

import argparse
import json
import sys
from datetime import date
from pathlib import Path
from typing import Any


_OUT_DIR = Path("./out")

# ---------------------------------------------------------------------------
# Excel raporu
# ---------------------------------------------------------------------------

def excel_raporu_uret(parseller: list[dict], dosya: str, baslik: str,
                      tarih: str) -> None:
    try:
        import openpyxl
        from openpyxl.styles import (
            Alignment, Border, Font, PatternFill, Side
        )
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("HATA: 'openpyxl' yüklü değil. pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.Workbook()

    # --- Özet sayfası ---
    ozet = wb.active
    ozet.title = "Özet"

    baslik_font = Font(name="Calibri", bold=True, color="FFFFFF", size=14)
    baslik_fill = PatternFill("solid", fgColor="1F4E79")
    alt_baslik_font = Font(name="Calibri", bold=True, size=11)
    koyu_border = Border(
        bottom=Side(border_style="medium", color="1F4E79")
    )

    ozet["A1"] = baslik
    ozet["A1"].font = baslik_font
    ozet["A1"].fill = baslik_fill
    ozet.merge_cells("A1:G1")
    ozet.row_dimensions[1].height = 28

    ozet["A2"] = f"Rapor Tarihi: {tarih}"
    ozet["A2"].font = Font(italic=True, color="595959")
    ozet["A3"] = f"Toplam Parsel: {len(parseller)}"
    ozet["A3"].font = alt_baslik_font
    toplam_alan = sum(
        float(p.get("alan_m2") or 0) for p in parseller if p.get("alan_m2")
    )
    ozet["A4"] = f"Toplam Alan: {toplam_alan:,.0f} m²"
    ozet["A4"].font = alt_baslik_font

    # Tablo başlıkları
    basliklar = [
        ("ADA", 10), ("PARSEL", 10), ("İL", 15), ("İLÇE", 18),
        ("MAHALLE", 22), ("ALAN (m²)", 14), ("NİTELİK", 16),
        ("PAFTA", 12), ("AÇIKLAMA", 30), ("HARİTA", 40),
    ]
    alan_sira = ["ada", "parsel", "il", "ilce", "mahalle", "alan_m2",
                 "nitelik", "pafta", "aciklama", "harita_url"]

    satir = 6
    for c, (bas, gen) in enumerate(basliklar, 1):
        hucre = ozet.cell(row=satir, column=c, value=bas)
        hucre.font = Font(bold=True, color="FFFFFF")
        hucre.fill = PatternFill("solid", fgColor="2E74B5")
        hucre.alignment = Alignment(horizontal="center", vertical="center")
        ozet.column_dimensions[get_column_letter(c)].width = gen

    for r, parsel in enumerate(parseller, satir + 1):
        dolu = (r - satir) % 2 == 1
        satir_fill = PatternFill("solid", fgColor="DEEAF1") if dolu else None
        for c, alan in enumerate(alan_sira, 1):
            deger = parsel.get(alan)
            if alan == "alan_m2" and deger is not None:
                try:
                    deger = float(deger)
                except (TypeError, ValueError):
                    pass
            hucre = ozet.cell(row=r, column=c, value=deger)
            if satir_fill:
                hucre.fill = satir_fill
            if alan == "alan_m2" and isinstance(deger, float):
                hucre.number_format = "#,##0.00"
            if alan == "harita_url" and deger:
                hucre.hyperlink = str(deger)
                hucre.value = "Haritaya Git"
                hucre.font = Font(color="0563C1", underline="single")

    # Hata sayfası (hatalı sorgular varsa)
    hatali = [p for p in parseller if p.get("hata")]
    if hatali:
        hata_ws = wb.create_sheet("Hatalar")
        hata_ws["A1"] = "SORGU HATALARI"
        hata_ws["A1"].font = Font(bold=True, color="FFFFFF")
        hata_ws["A1"].fill = PatternFill("solid", fgColor="C00000")
        hata_ws.merge_cells("A1:D1")
        for c, bas in enumerate(["ADA", "PARSEL", "İL", "HATA"], 1):
            hata_ws.cell(row=2, column=c, value=bas).font = Font(bold=True)
        for r, p in enumerate(hatali, 3):
            hata_ws.cell(row=r, column=1, value=p.get("ada"))
            hata_ws.cell(row=r, column=2, value=p.get("parsel"))
            hata_ws.cell(row=r, column=3, value=p.get("il"))
            hata_ws.cell(row=r, column=4, value=p.get("hata"))

    Path(dosya).parent.mkdir(parents=True, exist_ok=True)
    wb.save(dosya)
    print(f"Excel raporu kaydedildi: {dosya}")


# ---------------------------------------------------------------------------
# Word raporu
# ---------------------------------------------------------------------------

def word_raporu_uret(parseller: list[dict], dosya: str, baslik: str,
                     tarih: str) -> None:
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
    except ImportError:
        print("HATA: 'python-docx' yüklü değil. pip install python-docx", file=sys.stderr)
        sys.exit(1)

    doc = Document()

    # Sayfa kenar boşlukları
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1.2)

    # Başlık
    h = doc.add_heading(baslik, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    # Alt bilgi
    alt = doc.add_paragraph()
    alt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = alt.add_run(f"Rapor Tarihi: {tarih}  |  Toplam Parsel: {len(parseller)}")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    r.font.italic = True

    doc.add_paragraph()

    # Yönetici özeti
    doc.add_heading("Yönetici Özeti", level=1)
    toplam_alan = sum(
        float(p.get("alan_m2") or 0) for p in parseller if p.get("alan_m2")
    )
    il_sayisi = len({p.get("il", "") for p in parseller if p.get("il")})
    basarili = len([p for p in parseller if not p.get("hata")])

    ozet_metni = (
        f"Bu rapor {len(parseller)} adet taşınmaza ait parsel sorgu sonuçlarını "
        f"içermektedir. Sorgulama {tarih} tarihinde gerçekleştirilmiş olup "
        f"{basarili} parsel başarıyla sorgulanmıştır. "
        f"Toplam kadastral alan {toplam_alan:,.0f} m² ({toplam_alan / 10_000:,.2f} hektar) "
        f"olup {il_sayisi} ilde konuşlanmaktadır."
    )
    doc.add_paragraph(ozet_metni)
    doc.add_paragraph()

    # Parsel tablosu
    doc.add_heading("Parsel Detay Tablosu", level=1)

    sutunlar = ["Ada", "Parsel", "İl", "İlçe", "Mahalle", "Alan (m²)", "Nitelik"]
    alanlar = ["ada", "parsel", "il", "ilce", "mahalle", "alan_m2", "nitelik"]

    tablo = doc.add_table(rows=1, cols=len(sutunlar))
    tablo.style = "Table Grid"
    tablo.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Başlık satırı
    baslik_satir = tablo.rows[0]
    for c, (sut, alan) in enumerate(zip(sutunlar, alanlar)):
        hucre = baslik_satir.cells[c]
        hucre.text = sut
        hucre.paragraphs[0].runs[0].font.bold = True
        hucre.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # Arka plan rengi (OOXML)
        shading = hucre._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "2E74B5")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:val"), "clear")
        shading.append(shd)

    # Veri satırları
    for i, parsel in enumerate(parseller):
        satir = tablo.add_row()
        dolu = i % 2 == 0
        for c, alan in enumerate(alanlar):
            deger = parsel.get(alan)
            if alan == "alan_m2" and deger is not None:
                try:
                    deger = f"{float(deger):,.2f}"
                except (TypeError, ValueError):
                    deger = str(deger)
            hucre = satir.cells[c]
            hucre.text = str(deger or "")
            if dolu:
                shading = hucre._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "DEEAF1")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:val"), "clear")
                shading.append(shd)

    doc.add_paragraph()

    # Harita linkleri
    linkli = [p for p in parseller if p.get("harita_url")]
    if linkli:
        doc.add_heading("Harita Konumları", level=1)
        for p in linkli:
            kimlik = f"Ada {p.get('ada', '?')} / Parsel {p.get('parsel', '?')}"
            if p.get("aciklama"):
                kimlik += f" — {p['aciklama']}"
            para = doc.add_paragraph()
            para.add_run(f"{kimlik}: ").bold = True
            url_run = para.add_run(str(p["harita_url"]))
            url_run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            url_run.underline = True

    doc.add_paragraph()

    # Hatalı sorgular
    hatali = [p for p in parseller if p.get("hata")]
    if hatali:
        doc.add_heading("Sorgu Hataları", level=1)
        p = doc.add_paragraph()
        p.add_run("Aşağıdaki parseller için veri alınamadı:").bold = True
        for h in hatali:
            kimlik = f"Ada {h.get('ada', '?')}/Parsel {h.get('parsel', '?')}"
            doc.add_paragraph(
                f"{kimlik}: {h.get('hata', 'bilinmeyen hata')}",
                style="List Bullet",
            )

    # Notlar
    doc.add_paragraph()
    doc.add_heading("Notlar ve Yasal Uyarı", level=1)
    doc.add_paragraph(
        "Bu rapor TKGM Tapu ve Kadastro Bilgi Sistemi (TAKBİS) ve CBS "
        "web servislerinden otomatik olarak derlenen bilgileri içermektedir. "
        "Tapu sicilinin kesin durumu için yetkili Tapu Müdürlüğü'ne "
        "başvurulması gerekmektedir. Malik bilgileri ve tapu kütük şerhleri "
        "bu raporda yer almamaktadır.",
        style="Body Text",
    )

    Path(dosya).parent.mkdir(parents=True, exist_ok=True)
    doc.save(dosya)
    print(f"Word raporu kaydedildi: {dosya}")


# ---------------------------------------------------------------------------
# Ana akış
# ---------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description="Tapu/Parsel Rapor Üretici")
    ap.add_argument("--input", required=True,
                    help="Parsel sorgu sonuçları (.json)")
    ap.add_argument("--format", choices=["xlsx", "docx", "both"],
                    default="xlsx", help="Çıktı formatı")
    ap.add_argument("--baslik", default="Taşınmaz Analiz Raporu",
                    help="Rapor başlığı")
    ap.add_argument("--tarih", default=str(date.today()),
                    help="Rapor tarihi (YYYY-AA-GG)")
    ap.add_argument("--out-dir", default="./out",
                    help="Çıktı dizini (varsayılan: ./out)")
    args = ap.parse_args()

    with open(args.input, encoding="utf-8") as f:
        veriler: Any = json.load(f)

    # Tek obje veya liste kabul et
    if isinstance(veriler, dict):
        parseller = [veriler]
    elif isinstance(veriler, list):
        parseller = veriler
    else:
        print("HATA: JSON formatı geçersiz", file=sys.stderr)
        sys.exit(1)

    out = Path(args.out_dir)

    if args.format in ("xlsx", "both"):
        excel_raporu_uret(parseller, str(out / "tapu_raporu.xlsx"),
                          args.baslik, args.tarih)

    if args.format in ("docx", "both"):
        word_raporu_uret(parseller, str(out / "tapu_raporu.docx"),
                         args.baslik, args.tarih)


if __name__ == "__main__":
    main()
